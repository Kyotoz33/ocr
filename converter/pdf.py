from pathlib import Path

from .ocr import pil_to_rows, pil_to_text


def has_text(src: Path) -> bool:
    import pymupdf

    with pymupdf.open(src) as doc:
        return any(page.get_text().strip() for page in doc)


def render_pages(src: Path, dpi: int = 250):
    """Renderiza cada página do PDF como imagem PIL."""
    import pymupdf
    from PIL import Image

    with pymupdf.open(src) as doc:
        for page in doc:
            pix = page.get_pixmap(dpi=dpi)
            yield Image.frombytes("RGB", (pix.width, pix.height), pix.samples)


def pdf_to_docx(src: Path, dst: Path) -> None:
    """PDF -> DOCX. Usa OCR automaticamente se o PDF for escaneado (sem texto)."""
    if has_text(src):
        from pdf2docx import Converter

        cv = Converter(str(src))
        try:
            cv.convert(str(dst))
        finally:
            cv.close()
    else:
        _scanned_pdf_to_docx(src, dst)


def _scanned_pdf_to_docx(src: Path, dst: Path) -> None:
    from docx import Document

    doc = Document()
    for i, img in enumerate(render_pages(src)):
        if i:
            doc.add_page_break()
        for para in pil_to_text(img).split("\n\n"):
            para = " ".join(para.split())
            if para:
                doc.add_paragraph(para)
    doc.save(dst)


def pdf_to_xlsx(src: Path, dst: Path) -> None:
    """PDF -> XLSX: uma aba por tabela encontrada; OCR se o PDF for escaneado."""
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    if has_text(src):
        import pdfplumber

        with pdfplumber.open(src) as pdf:
            for pn, page in enumerate(pdf.pages, 1):
                for tn, table in enumerate(page.extract_tables(), 1):
                    ws = wb.create_sheet(f"p{pn}_t{tn}")
                    for row in table:
                        ws.append([c if c is not None else "" for c in row])
    else:
        for pn, img in enumerate(render_pages(src), 1):
            rows = pil_to_rows(img)
            if rows:
                ws = wb.create_sheet(f"p{pn}")
                for row in rows:
                    ws.append(row)
    if not wb.sheetnames:
        raise ValueError("Nenhuma tabela encontrada no PDF")
    wb.save(dst)
