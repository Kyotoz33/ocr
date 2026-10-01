import csv

import pytest
from PIL import Image, ImageDraw, ImageFont
from openpyxl import load_workbook

from converter import convert


def make_table_image(path):
    img = Image.new("RGB", (700, 200), "white")
    d = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 28)
    except OSError:
        font = ImageFont.load_default()
    for r, row in enumerate([("Nome", "Idade", "Cidade"), ("Ana", "30", "Recife"), ("Joao", "25", "Natal")]):
        for c, text in enumerate(row):
            d.text((20 + c * 230, 20 + r * 55), text, fill="black", font=font)
    img.save(path)


def test_image_to_csv_and_xlsx(tmp_path):
    img = tmp_path / "t.png"
    make_table_image(img)
    convert(img, tmp_path / "t.csv")
    rows = list(csv.reader(open(tmp_path / "t.csv", encoding="utf-8-sig")))
    assert rows[0] == ["Nome", "Idade", "Cidade"]
    assert rows[1][0] == "Ana"
    convert(img, tmp_path / "t.xlsx")
    ws = load_workbook(tmp_path / "t.xlsx").active
    assert ws["C3"].value == "Natal"


def test_pdf_to_docx(tmp_path):
    from reportlab.pdfgen import canvas

    pdf = tmp_path / "a.pdf"
    c = canvas.Canvas(str(pdf))
    c.drawString(100, 750, "Ola mundo")
    c.save()
    out = convert(pdf, tmp_path / "a.docx")
    from docx import Document

    assert "Ola mundo" in "\n".join(p.text for p in Document(out).paragraphs)


def test_unsupported(tmp_path):
    f = tmp_path / "x.txt"
    f.write_text("a")
    with pytest.raises(ValueError):
        convert(f, tmp_path / "x.docx")
